import re
import sys
from copy import deepcopy
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ENGLISH_PHRASES = [
    "capstone project",
    "e-commerce",
    "landing page",
    "user interface",
    "user experience",
    "use case",
    "sequence diagram",
    "activity diagram",
    "flowchart",
    "dashboard",
    "checkout",
    "template",
    "seller",
    "admin",
    "website",
    "marketplace",
    "responsive",
    "customization",
    "digital marketing",
    "framework",
    "platform",
    "login",
    "logout",
]
ENGLISH_WORDS = {
    "account", "admin", "analytics", "application", "backend", "banner", "brand", "browser", "business",
    "capstone", "cart", "catalog", "checkout", "client", "code", "content", "controller", "coupon",
    "css", "custom", "customer", "dashboard", "deploy", "design", "developer", "digital", "domain",
    "ecommerce", "endpoint", "feature", "file", "flowchart", "footer", "frontend", "framework", "hosting",
    "homepage", "html", "interface", "javascript", "landing", "layout", "login", "logout", "marketplace",
    "mobile", "navbar", "notification", "order", "page", "password", "platform", "product", "project",
    "register", "repository", "responsive", "review", "role", "schema", "seller", "sequence", "server",
    "session", "shop", "signin", "signup", "software", "source", "store", "support", "system", "template",
    "testing", "ui", "uml", "upload", "user", "ux", "web", "website", "workflow",
}


def set_font(run, size=12, bold=False, italic=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def clear_paragraph(paragraph):
    p = paragraph._element
    for child in list(p):
        if child.tag.endswith('}pPr'):
            continue
        p.remove(child)


def add_run(paragraph, text, size=12, bold=False, italic=False):
    run = paragraph.add_run(text)
    set_font(run, size=size, bold=bold, italic=italic)
    return run


def contains_drawing(paragraph):
    return bool(paragraph._element.xpath('.//w:drawing'))


def is_toc_entry(text):
    return ('\t' in text and re.search(r'\t\d+$', text.strip())) or bool(re.search(r'\.{3,}\s*\d+$', text.strip()))


def is_chapter_heading(text):
    clean = text.replace('\n', ' ').strip()
    return bool(re.match(r'^(BAB\s+([0-9]+|[IVX]+)|DAFTAR PUSTAKA|LAMPIRAN|ABSTRAK|ABSTRACT)\b', clean, re.I))


def is_subheading(text):
    clean = text.replace('\n', ' ').strip()
    return bool(re.match(r'^\d+(\.\d+){1,3}\s+', clean))


def is_table_title(text):
    return bool(re.match(r'^(Tabel|Table)\s+([0-9]+|[IVX]+|\d+\.\d+)', text.strip(), re.I))


def is_table_note(text):
    return bool(re.match(r'^(Sumber|Keterangan|Catatan)\s*[:.]', text.strip(), re.I))


def is_bibliography_entry(text):
    return bool(re.match(r'^[A-Z][A-Za-z\-\'`]+,', text.strip())) or bool(re.search(r'\(20\d{2}[a-z]?\)', text))


def iter_segments(text):
    lower = text.lower()
    matches = []
    for phrase in sorted(ENGLISH_PHRASES, key=len, reverse=True):
        for match in re.finditer(re.escape(phrase), lower):
            matches.append((match.start(), match.end(), text[match.start():match.end()], True))
    for match in re.finditer(r'\b[A-Za-z][A-Za-z\-]{2,}\b', text):
        word = match.group(0)
        if word.lower() in ENGLISH_WORDS:
            matches.append((match.start(), match.end(), word, True))
    matches.sort(key=lambda item: (item[0], -(item[1] - item[0])))

    filtered = []
    cursor = -1
    for item in matches:
        if item[0] >= cursor:
            filtered.append(item)
            cursor = item[1]

    out = []
    pos = 0
    for start, end, value, italic in filtered:
        if start > pos:
            out.append((text[pos:start], False))
        out.append((value, italic))
        pos = end
    if pos < len(text):
        out.append((text[pos:], False))
    return out or [(text, False)]


def format_paragraph_text(paragraph, text, size=12, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_indent=True, italicize_english=True):
    paragraph.alignment = align
    pf = paragraph.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.5
    pf.left_indent = Cm(0)
    pf.right_indent = Cm(0)
    pf.first_line_indent = Cm(1.25) if first_indent else Cm(0)

    clear_paragraph(paragraph)
    segments = iter_segments(text) if italicize_english else [(text, False)]
    for chunk, italic in segments:
        if chunk:
            add_run(paragraph, chunk, size=size, bold=bold, italic=italic)


def set_bibliography_format(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = paragraph.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.0
    pf.left_indent = Cm(0.75)
    pf.first_line_indent = Cm(-0.75)
    for run in paragraph.runs:
        set_font(run, size=12, bold=False, italic=run.italic)


def center_image_paragraphs(doc):
    for paragraph in doc.paragraphs:
        if contains_drawing(paragraph):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                if run.text:
                    set_font(run, size=12, italic=run.italic)


def resize_images(doc):
    for shape in doc.inline_shapes:
        max_width = Inches(5.8)
        if shape.width > max_width:
            ratio = float(shape.height) / float(shape.width)
            shape.width = int(max_width)
            shape.height = int(max_width * ratio)


def format_table(table):
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if row_index == 0 else WD_ALIGN_PARAGRAPH.JUSTIFY
                pf = paragraph.paragraph_format
                pf.space_before = Pt(0)
                pf.space_after = Pt(0)
                pf.line_spacing = 1.15
                pf.first_line_indent = Cm(0)
                for run in paragraph.runs:
                    set_font(run, size=12, bold=(row_index == 0), italic=run.italic)


def mark_table_borders(table):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement('w:tblBorders')
        tbl_pr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = 'w:' + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), '8')
        element.set(qn('w:space'), '0')
        element.set(qn('w:color'), '000000')


def iter_block_items(document):
    body = document.element.body
    for child in body.iterchildren():
        if child.tag.endswith('}p'):
            yield ('p', child)
        elif child.tag.endswith('}tbl'):
            yield ('tbl', child)


def main(input_path, output_path):
    doc = Document(input_path)
    paragraphs = doc.paragraphs
    tables = doc.tables

    para_by_el = {para._element: para for para in paragraphs}
    table_by_el = {table._element: table for table in tables}

    in_bibliography = False
    blocks = list(iter_block_items(doc))
    for idx, (kind, element) in enumerate(blocks):
        if kind == 'p':
            para = para_by_el[element]
            text = para.text.strip()
            if not text:
                continue

            clean = text.replace('\n', ' ').strip()

            if re.match(r'^LAMPIRAN\b', clean, re.I):
                in_bibliography = False

            if is_toc_entry(text):
                para.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for run in para.runs:
                    set_font(run, size=12, bold=False, italic=run.italic)
                continue

            if is_chapter_heading(text):
                in_bibliography = bool(re.match(r'^DAFTAR PUSTAKA\b', clean, re.I))
                format_paragraph_text(para, clean, size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=False, italicize_english=False)
                continue

            if is_subheading(text):
                format_paragraph_text(para, clean, size=12, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, first_indent=False, italicize_english=False)
                continue

            if is_table_title(text):
                format_paragraph_text(para, clean, size=12, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=False, italicize_english=False)
                continue

            if is_table_note(text):
                format_paragraph_text(para, clean, size=10, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, first_indent=False, italicize_english=False)
                continue

            if in_bibliography and is_bibliography_entry(text):
                set_bibliography_format(para)
                continue

            if contains_drawing(para):
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in para.runs:
                    set_font(run, size=12, bold=False, italic=run.italic)
                continue

            format_paragraph_text(para, clean, size=12, bold=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_indent=True, italicize_english=True)

        else:
            table = table_by_el[element]
            format_table(table)
            mark_table_borders(table)

    resize_images(doc)
    center_image_paragraphs(doc)
    for section in doc.sections:
        section.top_margin = Cm(4)
        section.left_margin = Cm(4)
        section.right_margin = Cm(3)
        section.bottom_margin = Cm(3)

    doc.save(output_path)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print("Usage: format_docx.py <input> <output>")
        sys.exit(1)
    main(sys.argv[1], sys.argv[2])
